#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Compare the engineer's blind labels with the current ground truth (registered experiment gt-blind-expert-v1); no answer is changed.
Main reading: agreement on the 84 scored cells (same quantity 36, no counterpart 48). Same quantity agrees only if the engineer also picks same
quantity with the same field, differencing and scale (within 2% relative); no counterpart agrees only if the engineer picks no counterpart.
Report only: three-class agreement and Cohen's kappa over all 285 cells, the labels given on different-definition cells, and every cell
labelled same quantity whose ground truth is not. Ground truth: realtel_adapter_v1/truth_v3.json (7 specifications, S01-S07) and
realtel_scale_v1/truth_final.json (12 client stacks, S08-S19). Reads a copy of the filled sheet whose sha256 must equal the original's.
Writes result.json and disagreements.md to experiments/results/gt_blind_expert_v1/. Usage: python3 scripts/paper/gt_blind_compare_v1.py
"""
import hashlib, json, re, shutil
from fractions import Fraction
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[2]; R = ROOT / "experiments/results"; OUT = R / "gt_blind_expert_v1"
SRC = ROOT / "Manual entry - ground truth blind labeling table - 2026-10-03.xlsx"
LAB = {"same quantity": "same", "different definition": "kind", "none": "none"}


def parse_scale(v):
    if v is None: return None
    t = str(v).strip().replace("×", "").replace("x", "").replace(" ", "").replace(",", "")
    if t in ("", "—", "-", "–"): return None
    try:
        if "/" in t: return float(Fraction(t))
        return float(t)
    except Exception:
        m = re.match(r"^1e([+-]?\d+)$", t, re.I)
        return 10.0 ** int(m.group(1)) if m else ("UNPARSED", str(v))


def parse_diff(v):
    t = (str(v).strip() if v is not None else "")
    return True if t in ("yes", "Y", "yes", "True") else False if t in ("no", "N", "no", "False") else None


def truth_lists():
    t1 = json.load(open(R / "realtel_adapter_v1/sources_test.json")); v3 = json.load(open(R / "realtel_adapter_v1/truth_v3.json"))["truth"]
    s2 = json.load(open(R / "realtel_scale_v1/sources.json")); tf = json.load(open(R / "realtel_scale_v1/truth_final.json"))["truth"]
    out = [(k, v3[k]) for k in t1] + [(n, tf[n]) for n in s2]
    assert len(out) == 19; return out


def kappa(a, b, cats=("same", "kind", "none")):
    n = len(a); po = sum(x == y for x, y in zip(a, b)) / n
    # Chance agreement from the two marginal label distributions.
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats); return (po - pe) / (1 - pe), po


def main():
    OUT.mkdir(parents=True, exist_ok=True); cp = OUT / "blind_sheet_filled.xlsx"
    if not cp.exists(): shutil.copy2(SRC, cp)
    h = hashlib.sha256(cp.read_bytes()).hexdigest(); assert h == hashlib.sha256(SRC.read_bytes()).hexdigest(), "the copy and the original do not match"
    wb = openpyxl.load_workbook(cp, data_only=True); T = truth_lists(); cells = []
    idx = {r[0]: r[1] for r in wb["source directory"].iter_rows(min_row=2, values_only=True) if r[0]}
    for i, (src, truth) in enumerate(T, 1):
        tab = f"S{i:02d}"; ws = wb[tab]; rows = list(ws.iter_rows(values_only=True))
        cand = [r[1] for r in rows if r and isinstance(r[0], str) and re.match(r"^F\d+$", r[0])]
        hdr = [j for j, r in enumerate(rows) if r and r[0] == "Feature index"][0]
        for k, r in enumerate(rows[hdr + 1: hdr + 16]):
            assert str(r[0]).strip() == str(k), (tab, r[0], k)
            lab = LAB.get(str(r[3]).strip() if r[3] else "", "missing"); fld = (str(r[4]).strip() if r[4] else None)
            t = truth[k]; c = {"page": tab, "page_name": idx.get(tab), "source": src, "k": k, "truth": t[0], "label": lab, "field": fld, "diff": parse_diff(r[5]), "scale": parse_scale(r[6]),
                               "basis": (str(r[7])[:400] if r[7] else None), "confidence": r[8], "field_in_candidates": (fld in cand) if fld else None}
            if t[0] == "same": c.update({"truth_field": t[1], "truth_diff": bool(t[2]), "truth_scale": t[3]})
            cells.append(c)
    assert len(cells) == 285
    def agree(c):
        if c["truth"] == "none": return c["label"] == "none"
        if c["truth"] == "same":
            sc = c["scale"]
            return (c["label"] == "same" and c["field"] == c["truth_field"] and c["diff"] == c["truth_diff"] and isinstance(sc, float) and abs(sc - c["truth_scale"]) <= 0.02 * abs(c["truth_scale"]))
        return None
    scored = [c for c in cells if c["truth"] in ("same", "none")]
    for c in scored: c["agree"] = agree(c)
    n_same = sum(c["truth"] == "same" for c in scored); n_none = sum(c["truth"] == "none" for c in scored)
    a_same = sum(c["agree"] for c in scored if c["truth"] == "same"); a_none = sum(c["agree"] for c in scored if c["truth"] == "none")
    kap, po = kappa([c["truth"] for c in cells], [c["label"] for c in cells])
    on_kind = {}
    for c in cells:
        if c["truth"] == "kind": on_kind[c["label"]] = on_kind.get(c["label"], 0) + 1
    extra_same = [c for c in cells if c["label"] == "same" and c["truth"] != "same"]
    same_label_wrong_detail = [c for c in scored if c["truth"] == "same" and not c["agree"]]
    none_disagree = [c for c in scored if c["truth"] == "none" and not c["agree"]]
    unparsed = [c for c in cells if isinstance(c["scale"], tuple)]
    res = {"sheet_sha256": h, "n_cells": 285, "scored": {"n": len(scored), "n_same": n_same, "n_none": n_none, "agree": a_same + a_none, "agree_rate": (a_same + a_none) / len(scored),
                                                          "agree_same": a_same, "agree_none": a_none},
           "three_class": {"agree_rate": po, "cohen_kappa": kap}, "labels_on_truth_kind": on_kind, "n_label_same_truth_not_same": len(extra_same), "unparsed_scales": len(unparsed),
           "cells": cells}
    (OUT / "result.json").write_text(json.dumps(res, ensure_ascii=False, indent=1, default=str))
    L = ["# Disagreements between the blind labels and the current answers (gt-blind-expert-v1, script-generated, to be checked item by item against the official original text)", "",
         f"Scoring cell {len(scored)} (same quantity {n_same}, none {n_none}) agreement {a_same + a_none}, agreement rate {(a_same + a_none) / len(scored):.3f}; 285 cells three-class agreement rate {po:.3f}, Cohen kappa {kap:.3f}.", "",
         "## 1. The current answer is the same quantity, the engineer disagrees"]
    for c in same_label_wrong_detail:
        L.append(f"- {c['page']} {c['page_name']} feature {c['k']}: current {c['truth_field']}, difference {c['truth_diff']}, coefficient {c['truth_scale']}; label {c['label']} {c['field']}, difference {c['diff']}, coefficient {c['scale']}, confidence {c['confidence']}. Basis: {c['basis']}")
    L += ["", "## 2. The current answer is none, the engineer disagrees"]
    for c in none_disagree:
        L.append(f"- {c['page']} {c['page_name']} feature {c['k']}: label {c['label']} {c['field'] or ''}, confidence {c['confidence']}. Basis: {c['basis']}")
    L += ["", "## 3. The engineer chose the same quantity but the current answer is not (including unscored cells)"]
    for c in extra_same:
        L.append(f"- {c['page']} {c['page_name']} feature {c['k']}: current {c['truth']}; label {c['field']}, difference {c['diff']}, coefficient {c['scale']}, confidence {c['confidence']}. Basis: {c['basis']}")
    (OUT / "disagreements.md").write_text("\n".join(L) + "\n")
    print(json.dumps({k: v for k, v in res.items() if k != "cells"}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
